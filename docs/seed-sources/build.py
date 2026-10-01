import csv
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

IN_F = "India|Bengaluru|Bangalore|Hyderabad|Pune|Chennai|Mumbai|Gurugram|Gurgaon|Noida|Delhi"
GH = "https://job-boards.greenhouse.io/"
LV = "https://jobs.lever.co/"
AS = "https://jobs.ashbyhq.com/"
SR = "https://careers.smartrecruiters.com/"

C, M = "search-confirmed", "from memory"
# (url, label, category, fit/evidence note, pay evidence, verified)
def gh(slug, label, cat, note, flt=None, pay="", v=C):
    return (GH + slug + (f"?location={flt}" if flt else ""), label, cat, note, pay, v)
def lv(slug, label, cat, note, flt=None, pay="", v=C):
    return (LV + slug + (f"?location={flt}" if flt else ""), label, cat, note, pay, v)
def ash(slug, label, cat, note, pay="", v=C):
    return (AS + slug, label, cat, note, pay, v)
def wd(host_site, label, cat, note, v=C):
    return ("https://" + host_site, label, cat, note, "", v)
def sr(slug, label, cat, note, v=C):
    return (SR + slug, label, cat, note, "", v)
def mem(url, label, cat, note="Not checked: careers page from memory"):
    return (url, label, cat, note, "", M)

ENTRY = "Entry/mid role seen: "
SENIOR = "Senior roles seen: "

india = [
 # --- Greenhouse (scraper supported today)
 gh("phonepe", "PhonePe", "Fintech", ENTRY + "Software Engineer Backend 1-3 yrs (Pune)"),
 gh("razorpaysoftwareprivatelimited", "Razorpay", "Fintech", "Lead SDE role seen; check SDE-1/2 openings"),
 gh("sigmoid", "Sigmoid", "Data engineering", ENTRY + "SDE I 1-3 yrs (Bengaluru), Java/Python/Scala"),
 gh("appdirect", "AppDirect", "SaaS", ENTRY + "SDE 1, 2+ yrs Java/Spring (Pune)", IN_F),
 gh("redwoodsoftware", "Redwood Software", "SaaS", ENTRY + "Associate SE 2+ yrs Java/Spring Boot (Hyderabad)", IN_F),
 gh("hyreo", "Hyreo", "SaaS", ENTRY + "Software Engineer Java 1-3 yrs"),
 gh("atomicwork", "Atomicwork", "SaaS startup", "Backend Engineer roles seen (Bengaluru)", IN_F),
 gh("tide", "Tide", "Fintech", "Senior backend seen; Java 17 / Spring Boot / jOOQ (Hyderabad)", IN_F),
 gh("agoda", "Agoda", "Travel tech", SENIOR + "Staff/Lead backend (Gurugram), Java/Kotlin/Scala", "Gurugram|Gurgaon|India"),
 gh("coupang", "Coupang", "E-commerce", SENIOR + "Staff backend (Bengaluru)", IN_F),
 gh("gleanwork", "Glean", "AI SaaS", "Backend Software Engineer seen (Bengaluru)", IN_F),
 gh("capco", "Capco", "Consulting (BFSI)", SENIOR + "Java backend 7+ yrs mostly (Bangalore/Pune/Chennai)", IN_F),
 gh("crashplan", "CrashPlan", "SaaS", SENIOR + "Senior Java 6+ yrs (Bengaluru)", IN_F),
 gh("addepar1", "Addepar", "Fintech", SENIOR + "Backend 5+ yrs (Pune)", IN_F),
 gh("pay2dc", "PayPay India", "Fintech", SENIOR + "Backend 4+ yrs, payments (Gurugram)"),
 gh("tegnaindia", "TEGNA India", "Media tech", SENIOR + "Lead Java 8+ yrs (Chennai)"),
 gh("blackduck", "Black Duck", "Security SaaS", "SDET 1 role seen", IN_F),
 gh("encora10", "Encora", "Engineering services", "Java developer roles seen", IN_F),
 gh("zetaglobal", "Zeta Global (marketing tech, not Zeta fintech)", "Martech", "AI/ML engineer seen (Bengaluru)", IN_F),
 # --- Lever (scraper supported today)
 lv("paytm", "Paytm", "Fintech", "Senior SE Backend 3-6 yrs Java/Spring Boot (Noida)"),
 lv("hevodata", "Hevo Data", "Data SaaS", ENTRY + "SDE I 0-1 yrs, Associate SDE (Bengaluru)", IN_F),
 lv("weekdayworks", "Weekday", "Recruiting platform", "Backend Engineer 3+ yrs Java/Spring Boot (Bengaluru)", None, "Rs 15-24 LPA listed"),
 lv("stable-money1", "Stable Money", "Fintech", ENTRY + "Software Engineer I Backend 1-2 yrs"),
 lv("fampay", "FamPay", "Fintech", SENIOR + "Senior SE 5-7 yrs (Bengaluru)"),
 lv("netomi", "Netomi", "AI SaaS", "Backend Engineer II 4+ yrs Java/Spring Boot (Gurugram)", IN_F),
 lv("yuno", "Yuno", "Payments", "Backend Java 3+ yrs, core payments (Hyderabad)", IN_F),
 lv("3pillarglobal", "3Pillar Global", "Engineering services", "Sr SE Java/Spring Boot/microservices 3-5 yrs", IN_F),
 lv("resilinc", "Resilinc", "Supply-chain SaaS", "Senior SE Java backend", IN_F),
 lv("acceldata", "Acceldata", "Data SaaS", "Bengaluru openings", IN_F),
 lv("zimperium", "Zimperium", "Security", "Java Engineer back-end & microservices (Bangalore)", IN_F),
 lv("gohighlevel", "HighLevel", "SaaS", "SDE II roles seen", IN_F),
 lv("jumpcloud", "JumpCloud", "SaaS", "Full-stack / devices SE India (Golang)", IN_F),
 lv("sonatype", "Sonatype", "DevSecOps", "Senior SE data, Java/Spring Batch (Hyderabad)", IN_F),
 lv("pditechnologies", "PDI Technologies", "SaaS", "Software Engineer III (Hyderabad/Chennai)", IN_F),
 lv("entrata", "Entrata", "SaaS", "Software Engineer (Pune)", IN_F),
 lv("everbridge", "Everbridge", "SaaS", "Senior SE backend Java/Spring Boot", IN_F),
 # --- Ashby (no scraper yet)
 ash("aiprise", "AiPrise", "Fintech/KYC", ENTRY + "SE I / II / III (Bangalore)"),
 ash("certifyos", "CertifyOS", "Healthtech", "SE Apps, Java or similar (Bengaluru)"),
 ash("bolna", "Bolna AI", "AI startup", "SE 2+ yrs (Bengaluru)"),
 ash("savvymoney", "SavvyMoney", "Fintech", SENIOR + "Sr SDE Backend 5+ yrs, remote India", "25-30 LPA listed"),
 ash("bjakcareer", "Bjak", "Fintech", "Backend / full-stack roles (India)"),
 ash("uipath", "UiPath", "Automation", SENIOR + "Java SE 5+ yrs"),
 ash("tekion", "Tekion", "Automotive SaaS", "SDET roles (Bangalore)"),
 ash("firstwork", "Firstwork", "HR tech", "SE 3-5 yrs, Python/Django"),
 ash("furtherai", "FurtherAI", "AI SaaS", SENIOR + "Backend/full-stack 5+ yrs (India)"),
 ash("cartesia", "Cartesia", "AI", SENIOR + "Platform SE 3-5 yrs (India)"),
 # --- Workday (no scraper yet)
 wd("mastercard.wd1.myworkdayjobs.com/CorporateCareers", "Mastercard", "Fintech / GCC", "Lead SE Java full stack (Pune)"),
 wd("paypal.wd1.myworkdayjobs.com/jobs", "PayPal", "Fintech / GCC", "Software Engineer Backend (Java) (Bangalore)"),
 wd("statestreet.wd1.myworkdayjobs.com/Global", "State Street", "BFSI / GCC", "Java full stack Sr Associate; Java backend (Bangalore/Hyderabad)"),
 wd("wf.wd1.myworkdayjobs.com/WellsFargoJobs", "Wells Fargo", "BFSI / GCC", "Sr SE Java full stack 4+ yrs (Hyderabad)"),
 wd("ms.wd5.myworkdayjobs.com/External", "Morgan Stanley", "BFSI / GCC", SENIOR + "Director-level Java (Mumbai/Bengaluru)"),
 wd("worldpay.wd5.myworkdayjobs.com/Worldpay_External_Careers_Site", "Worldpay", "Fintech", ENTRY + "Java Spring Boot 2-4 yrs (Pune/Bangalore/Indore)"),
 wd("jda.wd5.myworkdayjobs.com/JDA_Careers", "Blue Yonder (JDA)", "Supply-chain SaaS", "Staff SE Java/microservices/Spring Boot"),
 wd("redhat.wd5.myworkdayjobs.com/Jobs", "Red Hat", "Open source", ENTRY + "Associate SE trainee (Pune), Python-leaning"),
 wd("reliaquest.wd5.myworkdayjobs.com/ReliaQuest_Careers", "ReliaQuest", "Security", ENTRY + "Associate SE 0-1 yrs Java/Spring Boot (Pune)"),
 wd("cdk.wd1.myworkdayjobs.com/CDK", "CDK Global", "Automotive SaaS", ENTRY + "Associate SE Java/Spring (Hyderabad)"),
 wd("web.wd1.myworkdayjobs.com/ExternalCareerSite", "Newfold Digital / Bluehost", "Web hosting", "Java backend SE (Mumbai)"),
 wd("spgi.wd5.myworkdayjobs.com/spgi_internal", "S&P Global", "Fintech data", "Backend Java engineers (Hyderabad)"),
 wd("gartner.wd5.myworkdayjobs.com/EXT", "Gartner", "Research / GCC", "Senior SE Java backend primary"),
 wd("mavenir.wd1.myworkdayjobs.com/Mavenir_Careers", "Mavenir", "Telecom software", "Senior Java Developer (Bangalore)"),
 wd("motorolasolutions.wd5.myworkdayjobs.com/Careers", "Motorola Solutions", "Hardware + software", "SE Java/Spring Boot/microservices (Bangalore)"),
 wd("gevernova.wd5.myworkdayjobs.com/Vernova_ExternalSite", "GE Vernova", "Energy tech", "SE Java (Bengaluru)"),
 wd("hitachi.wd1.myworkdayjobs.com/hitachi", "Hitachi", "Conglomerate", "Java roles (Bengaluru)"),
 wd("hpe.wd5.myworkdayjobs.com/ACJobSite", "HPE", "Enterprise IT", "Engineering roles (Bengaluru)"),
 wd("nxp.wd3.myworkdayjobs.com/careers", "NXP Semiconductors", "Embedded", "Embedded firmware/software 3-5 yrs C/C++ (Hyderabad/Noida)"),
 wd("sensata.wd1.myworkdayjobs.com/Sensata-Careers", "Sensata", "Embedded", "Embedded Firmware Engineer (Pune)"),
 wd("hp.wd5.myworkdayjobs.com/ExternalCareerSite", "HP", "Embedded / hardware", "Embedded software / firmware roles"),
 # --- SmartRecruiters (no scraper yet)
 sr("servicenow?search=India", "ServiceNow", "SaaS", "India software roles (Hyderabad)"),
 sr("Visa", "Visa", "Fintech / GCC", "Software Engineer (Bangalore)"),
 sr("Experian", "Experian", "Fintech data / GCC", "Software Engineer Java + AWS 4+ yrs (Hyderabad)"),
 sr("Zscaler", "Zscaler", "Security", "Software Engineer II roles (Bengaluru)"),
 sr("WesternDigital", "Western Digital", "Embedded", ENTRY + "SDE Embedded, freshers batch (Bengaluru)"),
 sr("NECSWS", "NEC Software Solutions", "Software services", "Senior SE Java/Spring Boot/AWS/microservices (Mumbai)"),
 # --- From memory: product / startups
 mem("https://meesho.io/jobs", "Meesho", "E-commerce"),
 mem("https://groww.in/careers", "Groww", "Fintech"),
 mem("https://careers.cred.club/allJob", "CRED", "Fintech"),
 mem("https://careers.swiggy.com", "Swiggy", "Consumer tech"),
 mem("https://www.zomato.com/careers", "Zomato (Eternal)", "Consumer tech"),
 mem("https://www.flipkartcareers.com", "Flipkart", "E-commerce", "Not checked; SDE-1 ~21 LPA reported on a LeetCode thread (unverified)"),
 mem("https://www.zeta.tech/in/careers/", "Zeta (fintech)", "Fintech"),
 mem("https://juspay.io/careers", "Juspay", "Fintech"),
 mem("https://www.postman.com/company/careers/", "Postman", "SaaS"),
 mem("https://www.freshworks.com/company/careers/", "Freshworks", "SaaS"),
 mem("https://www.chargebee.com/company/careers/", "Chargebee", "SaaS"),
 mem("https://www.browserstack.com/careers", "BrowserStack", "SaaS"),
 mem("https://atlan.com/careers", "Atlan", "Data SaaS"),
 mem("https://clevertap.com/careers/", "CleverTap", "SaaS"),
 mem("https://www.moengage.com/careers/", "MoEngage", "SaaS"),
 mem("https://whatfix.com/careers/", "Whatfix", "SaaS"),
 mem("https://www.sprinklr.com/careers/", "Sprinklr", "SaaS"),
 mem("https://www.zoho.com/careers/", "Zoho", "SaaS"),
 mem("https://careers.makemytrip.com", "MakeMyTrip", "Travel tech"),
 mem("https://www.zeptonow.com/careers", "Zepto", "Consumer tech"),
 mem("https://www.olacabs.com/careers", "Ola", "Consumer tech"),
 mem("https://www.uber.com/global/en/careers/", "Uber", "Consumer tech"),
 mem("https://www.rubrik.com/company/careers", "Rubrik", "Infra software"),
 mem("https://www.nutanix.com/company/careers", "Nutanix", "Infra software"),
 mem("https://www.druva.com/about/careers", "Druva", "Infra software"),
 mem("https://www.cohesity.com/company/careers/", "Cohesity", "Infra software"),
 mem("https://www.thoughtspot.com/careers", "ThoughtSpot", "Analytics SaaS"),
 mem("https://careers.netapp.com", "NetApp", "Infra software / GCC"),
 mem("https://www.atlassian.com/company/careers", "Atlassian", "SaaS"),
 mem("https://careers.adobe.com", "Adobe", "Software / GCC"),
 mem("https://jobs.intuit.com", "Intuit", "Software / GCC"),
 mem("https://careers.salesforce.com", "Salesforce", "SaaS / GCC"),
 mem("https://careers.oracle.com", "Oracle", "Software / GCC"),
 mem("https://jobs.sap.com", "SAP", "Software / GCC"),
 mem("https://jobs.cisco.com", "Cisco", "Networking / GCC"),
 mem("https://jobs.dell.com", "Dell", "Hardware + software / GCC"),
 # --- From memory: GCC / BFSI / services with product teams
 mem("https://careers.walmart.com", "Walmart Global Tech", "Retail tech / GCC", "SE II roles (Bengaluru/Chennai) seen on aggregator; Java/C++ accepted"),
 mem("https://corporate.target.com/careers", "Target", "Retail tech / GCC"),
 mem("https://talent.lowes.com", "Lowe's", "Retail tech / GCC"),
 mem("https://www.americanexpress.com/en-in/careers/", "American Express", "Fintech / GCC"),
 mem("https://www.goldmansachs.com/careers/", "Goldman Sachs", "BFSI / GCC"),
 mem("https://careers.jpmorgan.com", "JPMorgan Chase", "BFSI / GCC"),
 mem("https://jobs.citi.com/location/india-jobs/287/1269750/2/1", "Citi India", "BFSI / GCC", "India software roles seen (Pune, Java/Python)"),
 mem("https://jobs.fidelity.com", "Fidelity Investments", "BFSI / GCC"),
 mem("https://careers.fiserv.com", "Fiserv", "Fintech"),
 mem("https://careers.fisglobal.com", "FIS", "Fintech"),
 mem("https://careers.broadridge.com", "Broadridge", "Fintech"),
 mem("https://careers.thomsonreuters.com", "Thomson Reuters", "Data / GCC"),
 mem("https://www.thoughtworks.com/careers", "Thoughtworks", "Consulting"),
 mem("https://www.nagarro.com/en/careers", "Nagarro", "Engineering services"),
 mem("https://careers.epam.com", "EPAM", "Engineering services"),
 mem("https://careers.publicissapient.com", "Publicis Sapient", "Consulting"),
 # --- From memory: embedded / hardware
 mem("https://nvidia.wd5.myworkdayjobs.com/NVIDIAExternalCareerSite", "NVIDIA", "Embedded / semis"),
 mem("https://cadence.wd1.myworkdayjobs.com/External_Careers", "Cadence", "Embedded / EDA"),
 mem("https://careers.ti.com", "Texas Instruments", "Embedded / semis"),
 mem("https://jobs.intel.com", "Intel", "Embedded / semis"),
 mem("https://careers.qualcomm.com", "Qualcomm", "Embedded / semis"),
 mem("https://careers.synopsys.com", "Synopsys", "Embedded / EDA"),
 mem("https://www.bosch.in/careers/", "Bosch", "Embedded / automotive"),
 mem("https://careers.honeywell.com", "Honeywell", "Embedded / industrial"),
 mem("https://jobs.siemens.com", "Siemens", "Embedded / industrial"),
 mem("https://www.careers.philips.com", "Philips", "Embedded / healthtech"),
 mem("https://jobs.harman.com", "Harman", "Embedded / automotive"),
 mem("https://www.continental.com/en/career/", "Continental", "Embedded / automotive"),
 mem("https://jobs.aptiv.com", "Aptiv", "Embedded / automotive"),
 mem("https://www.kpit.com/careers/", "KPIT", "Embedded / automotive"),
 mem("https://www.tataelxsi.com/careers", "Tata Elxsi", "Embedded / services"),
 mem("https://www.atherenergy.com/careers", "Ather Energy", "Embedded / EV"),
 mem("https://www.olaelectric.com/careers", "Ola Electric", "Embedded / EV"),
]
india_boards = [
 ("https://www.naukri.com", "Naukri", "Job board"), ("https://www.linkedin.com/jobs/", "LinkedIn Jobs", "Job board"),
 ("https://in.indeed.com", "Indeed India", "Job board"), ("https://www.foundit.in", "foundit (Monster India)", "Job board"),
 ("https://www.instahyre.com", "Instahyre", "Job board"), ("https://wellfound.com/jobs", "Wellfound (AngelList)", "Job board"),
 ("https://cutshort.io/jobs", "Cutshort", "Job board"), ("https://www.hirist.tech", "Hirist (tech jobs)", "Job board"),
 ("https://jobs.weekday.works", "Weekday", "Job board"), ("https://www.shine.com", "Shine", "Job board"),
 ("https://www.timesjobs.com", "TimesJobs", "Job board"),
]

uae = [
 gh("careem", "Careem", "Consumer tech", "Staff/Senior backend seen; Java/Golang; relocation to Dubai", "Dubai|UAE"),
 gh("cobblestoneenergy", "Cobblestone Energy", "Energy trading tech", "Software Engineer / Graduate SE (Dubai); visa + relocation", "Dubai|UAE", "USD 60-65k graduate SE listed (tax-free)"),
 gh("cobblestoneenergy4", "Cobblestone Energy (board 2)", "Energy trading tech", "Software Engineer / Graduate SE (Dubai)", "Dubai|UAE", "USD 60-65k graduate SE listed (tax-free)"),
 gh("propertyfinder", "Property Finder", "Proptech", "Uses Greenhouse; board slug unverified (first scrape will confirm)", "Dubai|UAE", v=M),
 lv("binance", "Binance", "Crypto", "Backend Engineer Java Web3 Wallet (Dubai)", "Dubai|UAE"),
 lv("1inch", "1inch", "Crypto", "Backend Software Engineer (Dubai)", "Dubai|UAE"),
 lv("palantir", "Palantir", "Software", SENIOR + "Forward Deployed SE (Abu Dhabi)", "Abu Dhabi|Dubai|UAE"),
 lv("Washmen", "Washmen", "Consumer tech", SENIOR + "Senior SE 7+ yrs (remote, UAE company)"),
 ash("syndica", "Syndica", "Crypto infra", "Rust/Go SE 3+ yrs (Dubai)"),
 sr("DeliveryHero?search=Dubai", "Delivery Hero (talabat)", "Consumer tech", ENTRY + "SE I Backend, SE Emirati-only, SE II (Dubai); .NET/Go/C# stacks"),
 sr("EtihadAirways5", "Etihad Airways", "Aviation", SENIOR + "Software Development Lead (Abu Dhabi)"),
 sr("masdar", "Masdar", "Energy", "Abu Dhabi roles"),
 sr("VAMS", "VAM Systems", "IT services (UAE)", "Java/MicroStrategy/other IT roles for UAE; sponsors visas"),
 sr("GhobashGroup", "Ghobash Group", "Conglomerate IT", "Data & AI / services engineer roles (Dubai)"),
 ("https://apply.workable.com/bayutdubizzle/", "Bayut & dubizzle", "Proptech / classifieds", "Technology roles; Workable board", "", C),
 ("https://careers.deliveryhero.com", "Delivery Hero careers (talabat jobs)", "Consumer tech", "SE II Backend, Eng Managers (Dubai)", "", C),
 ("https://www.emiratesnbd.com/en/careers", "Emirates NBD", "Banking", "Software/data/digital roles; many programmes are UAE-national only", "", C),
 mem("https://www.emiratesgroupcareers.com", "Emirates Group", "Aviation", "IT hiring drives reported (Java full stack, .NET, Python)"),
 mem("https://careers.adnoc.ae", "ADNOC", "Energy"),
 mem("https://www.bankfab.com/en-ae/about-fab/careers", "First Abu Dhabi Bank", "Banking"),
 mem("https://www.majidalfuttaim.com/en/careers", "Majid Al Futtaim", "Retail / tech"),
 mem("https://careers.microsoft.com", "Microsoft (UAE roles)", "Software"),
 mem("https://careers.oracle.com", "Oracle (UAE roles)", "Software"),
 mem("https://jobs.sap.com", "SAP (UAE roles)", "Software"),
 mem("https://jobs.cisco.com", "Cisco (UAE roles)", "Networking"),
 mem("https://www.ibm.com/careers", "IBM (UAE roles)", "Software"),
 mem("https://www.amazon.jobs/en/locations/united-arab-emirates", "Amazon / AWS UAE", "Software"),
]
uae_boards = [
 ("https://www.bayt.com/en/uae/jobs/", "Bayt UAE", "Job board"), ("https://www.naukrigulf.com", "Naukrigulf", "Job board"),
 ("https://ae.linkedin.com/jobs", "LinkedIn Jobs UAE", "Job board"), ("https://www.gulftalent.com/uae/jobs", "GulfTalent UAE", "Job board"),
 ("https://ae.indeed.com", "Indeed UAE", "Job board"), ("https://www.laimoon.com/uae", "Laimoon UAE", "Job board"),
 ("https://www.talentmate.com/jobs/uae", "TalentMate UAE", "Job board"), ("https://ae.mustakbil.com", "Mustakbil UAE", "Job board"),
 ("https://www.monstergulf.com", "Monster Gulf", "Job board"), ("https://relocate.me/united-arab-emirates", "Relocate.me UAE", "Job board (visa-sponsored)"),
 ("https://wellfound.com/jobs", "Wellfound (remote/UAE startups)", "Job board"),
]

def rows(entries, boards):
    out = []
    for url, label, cat, note, pay, ver in entries:
        out.append([url, label, True, cat, note, pay, ver])
    for url, label, cat in boards:
        out.append([url, label, True, cat, "Job board: read through email alerts (never scraped)", "", "known site"])
    return out

HEAD = ["url", "label", "enabled", "category", "fit_note", "pay_evidence", "verified"]
wb = Workbook()
for name, data in (("India", rows(india, india_boards)), ("UAE", rows(uae, uae_boards))):
    ws = wb.active if name == "India" else wb.create_sheet(name)
    ws.title = name
    ws.append(HEAD)
    for r in data:
        ws.append(r)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="DDEBF7")
    ws.freeze_panes = "A2"
    for i, w in enumerate([78, 38, 9, 24, 70, 34, 18], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    with open(f"sources_{name}.csv", "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f); w.writerow(HEAD); w.writerows(data)
wb.save("jobHunt_sources.xlsx")
print({"India": len(india) + len(india_boards), "UAE": len(uae) + len(uae_boards)})
